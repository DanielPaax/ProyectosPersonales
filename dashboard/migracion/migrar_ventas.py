#!/usr/bin/env python3
"""Migra la hoja 'BASE DE DATOS' del archivo de produccion a las hojas VENTAS y
TARIFAS del esquema 3. Escribe VALORES, no formulas: los costos quedan
congelados al momento de la venta.

Uso:  python3 migrar_ventas.py ARCHIVO_DE_PRODUCCION_2026.xlsx [salida.xlsx]
Requiere: pip install openpyxl
"""
import sys, re, datetime, collections, warnings
import openpyxl

warnings.filterwarnings('ignore')

ENCABEZADOS_VENTAS = [
    'ID', 'FECHA_VENTA', 'SEM_INICIO', 'SEMANA', 'FECHA_PROD', 'SEMANA_ORIG',
    'ODC', 'ODP', 'CLIENTE', 'TIPO_PROYECTO', 'PROYECTO', 'MAQUINA',
    'MATERIAL', 'M2_PROYECTO', 'M2_REALES', 'VENTA', 'COSTO_MATERIAL',
    'COSTO_MERMA', 'MAQUILA', 'INSTALACION', 'MANO_OBRA', 'HORAS_EXTRAS',
    'OTROS', 'CARGA_ADMIN', 'COMERCIAL', 'FACTURA', 'FECHA_FACTURA',
    'METODO_PAGO', 'ESTATUS_COBRO', 'ESTATUS_MONTO', 'CELDA_ORIGEN',
    'OBS_MIGRACION', 'VENDEDOR'
]
COLS = {n: i for i, n in enumerate(ENCABEZADOS_VENTAS)}

# Unicos cambios de nombre automaticos: los deterministas. Lo dudoso se reporta.
ALIAS_CLIENTE = {'FIBRA SHOP': 'FIBRASHOP'}


def txt(v):
    return re.sub(r'\s+', ' ', str(v)).strip() if v not in (None, '') else ''


def num(v):
    """Numero o None. Los '#N/A' y '#DIV/0!' de Excel no son numeros."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def lunes(d):
    return d - datetime.timedelta(days=d.weekday())


def etiqueta_semana(l):
    fin = l + datetime.timedelta(days=5)
    return 'SEM %s-%s' % (l.strftime('%d/%m'), fin.strftime('%d/%m/%y'))


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    ruta = sys.argv[1]
    salida = sys.argv[2] if len(sys.argv) > 2 else 'VENTAS_migracion.xlsx'
    hoy = datetime.datetime.now()

    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb['BASE DE DATOS']
    lst = wb['LISTAS']

    # ---- tarifas: costo por m2 de cada material y carga por m2 -------------
    tarifas = []
    for r in lst.iter_rows(min_row=3, values_only=True):
        if r[8] and num(r[9]) is not None:
            tarifas.append((txt(r[8]), round(num(r[9]), 4)))
    carga_m2 = num(lst['H3'].value)

    filas = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if r[1] is None:          # sin fecha de venta = renglon vacio con #N/A
            continue
        filas.append((i, r))

    # monto por pedido, para saber si a un renglon sin monto le "cubre" otro
    def pedido(r):
        return txt(r[4]) or txt(r[5])
    monto_pedido = collections.defaultdict(float)
    for _, r in filas:
        v = num(r[29])
        if v:
            monto_pedido[pedido(r)] += v

    out, resumen = [], collections.Counter()
    clientes = collections.Counter()
    for n, (xl, r) in enumerate(filas, start=1):
        fv, fp = r[1], r[2]
        l = lunes(fv.date())
        obs = []

        cliente = txt(r[6]).upper()
        cliente = ALIAS_CLIENTE.get(cliente, cliente)
        clientes[cliente] += 1

        venta = num(r[29])
        pendiente = not venta
        if pendiente:
            if pedido(r) and monto_pedido[pedido(r)] > 0:
                obs.append('Sin monto propio; otra linea del mismo pedido (%s) '
                           'si trae monto: revisar si ya esta incluido ahi.' % pedido(r))
                resumen['pend_con_pista'] += 1
            else:
                obs.append('Pedido sin monto de venta en ninguna linea.')
                resumen['pend_sin_pista'] += 1

        merma, maq, inst = num(r[16]) or 0, num(r[18]) or 0, num(r[19]) or 0
        mo, hext, otros = num(r[20]) or 0, num(r[21]) or 0, num(r[22]) or 0
        carga = num(r[27]) or 0
        x = num(r[23])
        if x is None:
            c_mat = 0.0
            obs.append('Costo no calculable en el archivo original (material sin tarifa).')
            resumen['costo_na'] += 1
        else:
            c_mat = x - (merma + maq + inst + mo + hext + otros)

        if fv > hoy:
            obs.append('Fecha de venta futura (%s).' % fv.strftime('%d/%m/%Y'))
            resumen['fecha_futura'] += 1
        if fp and abs((fv - fp).days) > 90:
            obs.append('Fecha de venta y de produccion difieren %d dias.' % abs((fv - fp).days))
            resumen['fechas_distantes'] += 1

        factura = txt(r[34])
        out.append([
            n, fv, l, etiqueta_semana(l), fp, txt(r[3]),
            txt(r[4]), txt(r[5]), cliente, txt(r[8]), txt(r[10]),
            txt(r[11]) or '(NO REPORTADA)', txt(r[12]) or '(NO REPORTADO)',
            num(r[13]), num(r[14]),
            round(venta, 2) if venta else None,
            round(c_mat, 4), round(merma, 4), round(maq, 4), round(inst, 4),
            round(mo, 4), round(hext, 4), round(otros, 4), round(carga, 4),
            txt(r[32]).upper(), factura, r[35] if r[35] else None,
            txt(r[33]), 'FACTURADO' if factura else '',
            'PENDIENTE' if pendiente else 'CON MONTO',
            "BASE DE DATOS!B%d" % xl, ' '.join(obs), ''
        ])

    # ---- conciliacion contra el archivo original ----------------------------
    v_orig = sum(num(r[29]) or 0 for _, r in filas)
    c_orig = sum(num(r[28]) or 0 for _, r in filas)
    v_new = sum(f[COLS['VENTA']] or 0 for f in out)
    c_new = sum(sum(f[COLS['COSTO_MATERIAL']:COLS['CARGA_ADMIN'] + 1]) for f in out)
    ok_v = abs(v_orig - v_new) < 0.5
    ok_c = abs(c_orig - c_new) < 1.0

    # ---- libro de salida -----------------------------------------------------
    nuevo = openpyxl.Workbook()
    v = nuevo.active
    v.title = 'VENTAS'
    v.append(ENCABEZADOS_VENTAS)
    for f in out:
        v.append(f)
    for fila in v.iter_rows(min_row=2):
        for idx in (COLS['FECHA_VENTA'], COLS['SEM_INICIO'], COLS['FECHA_PROD'], COLS['FECHA_FACTURA']):
            fila[idx].number_format = 'dd/mm/yyyy'
    v.freeze_panes = 'A2'

    t = nuevo.create_sheet('TARIFAS')
    t.append(['MATERIAL', 'COSTO_M2', None, 'PARAMETRO', 'VALOR'])
    for k, (m, c) in enumerate(tarifas, start=2):
        t.cell(k, 1, m)
        t.cell(k, 2, c)
    t.cell(2, 4, 'CARGA_M2')
    t.cell(2, 5, carga_m2)
    t.freeze_panes = 'A2'
    nuevo.save(salida)

    # ---- reporte --------------------------------------------------------------
    print('Renglones migrados: %d (descartados vacios: %d)' %
          (len(out), ws.max_row - 1 - len(out)))
    print('Venta   original %.2f | migrada %.2f | %s' % (v_orig, v_new, 'OK' if ok_v else 'NO CUADRA'))
    print('Costo   original %.2f | migrado %.2f | %s' % (c_orig, c_new, 'OK' if ok_c else 'NO CUADRA'))
    print('Pendientes de monto: %d  (con pista de otra linea: %d | pedido sin monto: %d)' %
          (resumen['pend_con_pista'] + resumen['pend_sin_pista'],
           resumen['pend_con_pista'], resumen['pend_sin_pista']))
    print('Fechas de venta futuras: %d | fechas venta/prod distantes (>90d): %d | costo no calculable: %d' %
          (resumen['fecha_futura'], resumen['fechas_distantes'], resumen['costo_na']))
    print('Tarifas de material: %d | carga por m2: %s' % (len(tarifas), carga_m2))
    print('Clientes tras normalizar (%d):' % len(clientes))
    for c, k in sorted(clientes.items()):
        print('   %-32s %d' % (c, k))
    print('Salida:', salida)
    if not (ok_v and ok_c):
        sys.exit(2)


if __name__ == '__main__':
    main()
